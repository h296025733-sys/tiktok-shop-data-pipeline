from __future__ import annotations

import csv
import hashlib
import json
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from src.config import PROJECT_ROOT
from src.xlsx_utils import write_rows_to_workbook


FORMULA_ERROR_TOKENS = ("#REF!", "#DIV/0!", "#N/A", "#VALUE!", "#NAME?", "#NUM!")
SHEET_REFERENCE_RE = re.compile(
    r"(?:'((?:[^']|'')+)'|([A-Za-z0-9_.\-\u3400-\u9fff]+))!"
)
RAW_SNAPSHOT_NAMES = {"6.17", "1.21", "12.23"}
SENSITIVE_HEADER_TERMS = {
    "邮箱",
    "whatsapp",
    "联系方式",
    "付款渠道",
    "付款账号",
    "paypal",
    "备注",
}


@dataclass
class SheetStats:
    title: str
    state: str
    declared_max_row: int
    declared_max_column: int
    data_min_row: int | None
    data_max_row: int | None
    data_min_column: int | None
    data_max_column: int | None
    non_empty_cells: int
    formula_cells: int
    literal_cells: int
    styled_blank_cells: int
    merged_ranges: int
    data_validations: int
    hyperlinks: int


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def _safe_sample(value: Any, limit: int = 120) -> str:
    text = _text(value).replace("\r", " ").replace("\n", " ")
    return text if len(text) <= limit else f"{text[: limit - 1]}…"


def _source_assessment(sheet: str, column_index: int) -> tuple[str, str]:
    if sheet in RAW_SNAPSHOT_NAMES:
        return "TikTok Seller Center 手工导出", "平台原始字段（直接）"
    if sheet == "Tech-wave红人开发表":
        return "公司内部录入", "公司人工 CRM"
    if sheet == "Tech-wave红人合作表":
        if column_index in {4, 5, 9}:
            return "公式/程序计算与公司数据", "混合：计算字段或内部人工字段"
        return "公司内部录入", "公司人工 CRM"
    if sheet == "Tech-wave红人上线表":
        if column_index in {12, 13, 14}:
            return "日期快照公式回填", "TikTok 手工报表字段（直接/断链）"
        if column_index in {1, 2, 3, 8, 20}:
            return "公式或合作表带入", "程序计算/间接依赖"
        return "公司内部录入或待确认", "混合/待确认"
    if sheet == "投产比":
        return "公式、人工值及多外部报表混合", "经营汇总混合字段"
    return "非主数据链", "说明/建议"


def _formula_references(formula: str) -> list[str]:
    references: list[str] = []
    for match in SHEET_REFERENCE_RE.finditer(formula):
        name = (match.group(1) or match.group(2) or "").replace("''", "'")
        if name and name.upper() != "REF":
            references.append(name)
    return sorted(set(references))


def _header_for_column(worksheet, column_index: int) -> tuple[int, str]:
    header_rows = range(1, 4) if worksheet.title == "投产比" else range(1, 2)
    parts = []
    for row_index in header_rows:
        value = _text(worksheet.cell(row_index, column_index).value)
        if value:
            parts.append(value)
    header = " | ".join(dict.fromkeys(parts))
    if not header and worksheet.title in RAW_SNAPSHOT_NAMES and column_index == 1:
        header = "视频链接（原表头为空，按数据内容识别）"
    return (3 if worksheet.title == "投产比" else 1), header


def _write_csv(path: Path, headers: list[str], rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def audit_workbook(source_path: Path) -> dict[str, Any]:
    source_path = source_path.resolve()
    workbook = load_workbook(source_path, read_only=False, data_only=False, keep_links=True)

    sheet_stats: list[SheetStats] = []
    field_rows: list[dict[str, Any]] = []
    dependency_rows: list[dict[str, Any]] = []
    error_rows: list[dict[str, Any]] = []
    relation_counts: Counter[tuple[str, str]] = Counter()
    ref_error_counts: Counter[str] = Counter()
    formula_counts: Counter[str] = Counter()
    formula_ref_snapshot_counts: Counter[tuple[str, int]] = Counter()
    formula_ref_other_counts: Counter[tuple[str, int]] = Counter()
    formula_error_column_counts: Counter[tuple[str, int]] = Counter()

    for worksheet in workbook.worksheets:
        non_empty = []
        formula_count = 0
        literal_count = 0
        styled_blank_count = 0
        hyperlink_count = 0
        column_counts: dict[int, Counter[str]] = defaultdict(Counter)
        column_formats: dict[int, Counter[str]] = defaultdict(Counter)
        column_samples: dict[int, list[str]] = defaultdict(list)
        column_first_row: dict[int, int] = {}
        column_last_row: dict[int, int] = {}

        for cell in worksheet._cells.values():
            if cell.hyperlink:
                hyperlink_count += 1
            if cell.value is None:
                if cell.has_style or cell.comment or cell.hyperlink:
                    styled_blank_count += 1
                continue

            non_empty.append(cell)
            column_first_row[cell.column] = min(
                column_first_row.get(cell.column, cell.row), cell.row
            )
            column_last_row[cell.column] = max(
                column_last_row.get(cell.column, cell.row), cell.row
            )
            column_formats[cell.column][cell.number_format] += 1
            if len(column_samples[cell.column]) < 3 and cell.row > 1:
                column_samples[cell.column].append(_safe_sample(cell.value))

            if cell.data_type == "f" or (
                isinstance(cell.value, str) and cell.value.startswith("=")
            ):
                formula_count += 1
                formula_counts[worksheet.title] += 1
                column_counts[cell.column]["formula"] += 1
                formula = str(cell.value)
                references = _formula_references(formula)
                for target_sheet in references:
                    relation_counts[(worksheet.title, target_sheet)] += 1
                if any(target in RAW_SNAPSHOT_NAMES for target in references):
                    formula_ref_snapshot_counts[(worksheet.title, cell.column)] += 1
                if any(target not in RAW_SNAPSHOT_NAMES for target in references):
                    formula_ref_other_counts[(worksheet.title, cell.column)] += 1
                tokens = [token for token in FORMULA_ERROR_TOKENS if token in formula]
                formula_hash = hashlib.sha256(formula.encode("utf-8")).hexdigest()
                dependency_rows.append(
                    {
                        "source_sheet": worksheet.title,
                        "source_cell": cell.coordinate,
                        "referenced_sheets": "|".join(references),
                        "reference_count": len(references),
                        "contains_formula_error": bool(tokens),
                        "formula_sha256": formula_hash,
                        "formula": formula,
                    }
                )
                if tokens:
                    ref_error_counts[worksheet.title] += 1
                    formula_error_column_counts[(worksheet.title, cell.column)] += 1
                    error_rows.append(
                        {
                            "sheet": worksheet.title,
                            "cell": cell.coordinate,
                            "error_source": "formula_text",
                            "error_tokens": "|".join(tokens),
                            "formula_or_value": formula,
                        }
                    )
            else:
                literal_count += 1
                column_counts[cell.column]["literal"] += 1

        data_rows = [cell.row for cell in non_empty]
        data_columns = [cell.column for cell in non_empty]
        stats = SheetStats(
            title=worksheet.title,
            state=worksheet.sheet_state,
            declared_max_row=worksheet.max_row,
            declared_max_column=worksheet.max_column,
            data_min_row=min(data_rows) if data_rows else None,
            data_max_row=max(data_rows) if data_rows else None,
            data_min_column=min(data_columns) if data_columns else None,
            data_max_column=max(data_columns) if data_columns else None,
            non_empty_cells=len(non_empty),
            formula_cells=formula_count,
            literal_cells=literal_count,
            styled_blank_cells=styled_blank_count,
            merged_ranges=len(worksheet.merged_cells.ranges),
            data_validations=len(worksheet.data_validations.dataValidation),
            hyperlinks=hyperlink_count,
        )
        sheet_stats.append(stats)

        for column_index in range(1, worksheet.max_column + 1):
            counts = column_counts[column_index]
            non_empty_count = counts["formula"] + counts["literal"]
            if non_empty_count == 0 and column_index not in column_first_row:
                continue
            header_row, header = _header_for_column(worksheet, column_index)
            formula_count_column = counts["formula"]
            literal_count_column = counts["literal"]
            direct_formula_count = formula_ref_snapshot_counts[
                (worksheet.title, column_index)
            ]
            other_formula_count = formula_ref_other_counts[(worksheet.title, column_index)]
            tiktok_direct_count = (
                max(non_empty_count - 1, 0) if worksheet.title in RAW_SNAPSHOT_NAMES else 0
            )
            source, owner = _source_assessment(worksheet.title, column_index)
            sensitive_header = any(
                term in header.lower() for term in SENSITIVE_HEADER_TERMS
            )
            field_rows.append(
                {
                    "sheet": worksheet.title,
                    "column": get_column_letter(column_index),
                    "column_index": column_index,
                    "header_row": header_row,
                    "header": header,
                    "first_non_empty_row": column_first_row.get(column_index, ""),
                    "last_non_empty_row": column_last_row.get(column_index, ""),
                    "non_empty_cells": non_empty_count,
                    "formula_cells": formula_count_column,
                    "literal_cells": literal_count_column,
                    "formula_ratio": round(
                        formula_count_column / non_empty_count, 6
                    )
                    if non_empty_count
                    else 0,
                    "literal_ratio": round(
                        literal_count_column / non_empty_count, 6
                    )
                    if non_empty_count
                    else 0,
                    "direct_snapshot_formula_cells": direct_formula_count,
                    "other_cross_sheet_formula_cells": other_formula_count,
                    "formula_error_cells": formula_error_column_counts[
                        (worksheet.title, column_index)
                    ],
                    "tiktok_direct_cells_estimate": tiktok_direct_count,
                    "source_assessment": source,
                    "field_owner": owner,
                    "number_formats": "|".join(
                        item[0] for item in column_formats[column_index].most_common(5)
                    ),
                    "samples": (
                        "已抑制敏感字段样例"
                        if sensitive_header
                        else " | ".join(column_samples[column_index])
                    ),
                }
            )

    # Cached formula errors are read separately so formula text and last calculated values are not mixed.
    cached_workbook = load_workbook(
        source_path, read_only=True, data_only=True, keep_links=False
    )
    existing_error_cells = {(row["sheet"], row["cell"]) for row in error_rows}
    for worksheet in cached_workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                value = cell.value
                if (
                    isinstance(value, str)
                    and value in FORMULA_ERROR_TOKENS
                    and (worksheet.title, cell.coordinate) not in existing_error_cells
                ):
                    error_rows.append(
                        {
                            "sheet": worksheet.title,
                            "cell": cell.coordinate,
                            "error_source": "cached_value",
                            "error_tokens": value,
                            "formula_or_value": value,
                        }
                    )
    cached_workbook.close()

    report_dir = PROJECT_ROOT / "exports" / "reports"
    sheet_inventory_path = report_dir / "workbook_sheet_inventory.csv"
    field_inventory_csv = report_dir / "workbook_field_inventory.csv"
    field_inventory_xlsx = report_dir / "workbook_field_inventory.xlsx"
    dependencies_path = report_dir / "workbook_formula_dependencies.csv"
    errors_path = report_dir / "workbook_formula_errors.csv"
    relationship_xlsx = report_dir / "workbook_relationship_matrix.xlsx"
    relationship_csv = report_dir / "workbook_relationship_matrix.csv"

    sheet_headers = [
        "sheet",
        "state",
        "declared_max_row",
        "declared_max_column",
        "data_min_row",
        "data_max_row",
        "data_min_column",
        "data_max_column",
        "non_empty_cells",
        "formula_cells",
        "literal_cells",
        "styled_blank_cells",
        "merged_ranges",
        "data_validations",
        "hyperlinks",
    ]
    sheet_rows = [
        {
            "sheet": item.title,
            "state": item.state,
            "declared_max_row": item.declared_max_row,
            "declared_max_column": item.declared_max_column,
            "data_min_row": item.data_min_row,
            "data_max_row": item.data_max_row,
            "data_min_column": item.data_min_column,
            "data_max_column": item.data_max_column,
            "non_empty_cells": item.non_empty_cells,
            "formula_cells": item.formula_cells,
            "literal_cells": item.literal_cells,
            "styled_blank_cells": item.styled_blank_cells,
            "merged_ranges": item.merged_ranges,
            "data_validations": item.data_validations,
            "hyperlinks": item.hyperlinks,
        }
        for item in sheet_stats
    ]
    _write_csv(sheet_inventory_path, sheet_headers, sheet_rows)

    field_headers = list(field_rows[0].keys()) if field_rows else []
    _write_csv(field_inventory_csv, field_headers, field_rows)
    write_rows_to_workbook(
        field_inventory_xlsx,
        sheet_name="字段盘点",
        headers=field_headers,
        rows=([row[header] for header in field_headers] for row in field_rows),
        table_name="FieldInventory",
    )

    dependency_headers = [
        "source_sheet",
        "source_cell",
        "referenced_sheets",
        "reference_count",
        "contains_formula_error",
        "formula_sha256",
        "formula",
    ]
    _write_csv(dependencies_path, dependency_headers, dependency_rows)

    error_headers = [
        "sheet",
        "cell",
        "error_source",
        "error_tokens",
        "formula_or_value",
    ]
    _write_csv(errors_path, error_headers, error_rows)

    relationship_rows = []
    for (source_sheet, target_sheet), count in sorted(
        relation_counts.items(), key=lambda item: (-item[1], item[0])
    ):
        relationship_rows.append(
            {
                "source_sheet": source_sheet,
                "target_sheet": target_sheet,
                "formula_cell_count": count,
                "target_exists": target_sheet in workbook.sheetnames,
                "relationship": (
                    "日期快照直接回填"
                    if target_sheet in RAW_SNAPSHOT_NAMES
                    else "跨业务表公式依赖"
                ),
            }
        )
    relationship_headers = list(relationship_rows[0].keys()) if relationship_rows else [
        "source_sheet",
        "target_sheet",
        "formula_cell_count",
        "target_exists",
        "relationship",
    ]
    _write_csv(relationship_csv, relationship_headers, relationship_rows)
    write_rows_to_workbook(
        relationship_xlsx,
        sheet_name="关系矩阵",
        headers=relationship_headers,
        rows=(
            [row[header] for header in relationship_headers]
            for row in relationship_rows
        ),
        table_name="RelationshipMatrix",
    )

    snapshot_summary = _snapshot_summary(workbook["6.17"])
    external_link_count = len(getattr(workbook, "_external_links", []))
    workbook.close()

    diagnosis_path = PROJECT_ROOT / "docs" / "01_workbook_diagnosis.md"
    diagnosis_path.write_text(
        _diagnosis_markdown(
            source_path=source_path,
            sheet_stats=sheet_stats,
            snapshot_summary=snapshot_summary,
            formula_counts=formula_counts,
            ref_error_counts=ref_error_counts,
            formula_error_column_counts=formula_error_column_counts,
            error_rows=error_rows,
            relationship_rows=relationship_rows,
            external_link_count=external_link_count,
        ),
        encoding="utf-8",
    )

    summary = {
        "source_file": str(source_path),
        "sheet_count": len(sheet_stats),
        "sheet_names": [item.title for item in sheet_stats],
        "hidden_sheets": [
            item.title for item in sheet_stats if item.state != "visible"
        ],
        "formula_cells": sum(formula_counts.values()),
        "formula_error_rows": len(error_rows),
        "formula_text_with_ref_error": sum(ref_error_counts.values()),
        "external_links": external_link_count,
        "snapshot_6_17": snapshot_summary,
        "outputs": {
            "diagnosis": str(diagnosis_path),
            "sheet_inventory": str(sheet_inventory_path),
            "field_inventory": str(field_inventory_xlsx),
            "formula_dependencies": str(dependencies_path),
            "formula_errors": str(errors_path),
            "relationship_matrix": str(relationship_xlsx),
        },
    }
    (report_dir / "workbook_audit_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return summary


def _snapshot_summary(worksheet) -> dict[str, Any]:
    rows = []
    for row_index in range(2, worksheet.max_row + 1):
        values = [worksheet.cell(row_index, column).value for column in range(1, 17)]
        if any(value is not None for value in values):
            rows.append(values)
    urls = [_text(row[0]) for row in rows if _text(row[0])]
    creator_ids = [_text(row[2]) for row in rows if _text(row[2])]
    video_ids = [_text(row[4]) for row in rows if _text(row[4])]
    publish_dates = [
        row[5].date() if isinstance(row[5], datetime) else row[5]
        for row in rows
        if isinstance(row[5], (datetime, date))
    ]
    full_rows = [json.dumps([_text(value) for value in row], ensure_ascii=False) for row in rows]
    return {
        "data_rows": len(rows),
        "distinct_video_urls": len(set(urls)),
        "duplicate_video_url_rows": len(urls) - len(set(urls)),
        "distinct_video_ids": len(set(video_ids)),
        "distinct_creator_ids": len(set(creator_ids)),
        "fully_duplicate_rows": len(full_rows) - len(set(full_rows)),
        "publish_date_min": min(publish_dates).isoformat() if publish_dates else None,
        "publish_date_max": max(publish_dates).isoformat() if publish_dates else None,
        "headers": [_text(worksheet.cell(1, column).value) for column in range(1, 17)],
    }


def _diagnosis_markdown(
    *,
    source_path: Path,
    sheet_stats: list[SheetStats],
    snapshot_summary: dict[str, Any],
    formula_counts: Counter[str],
    ref_error_counts: Counter[str],
    formula_error_column_counts: Counter[tuple[str, int]],
    error_rows: list[dict[str, Any]],
    relationship_rows: list[dict[str, Any]],
    external_link_count: int,
) -> str:
    hidden = [item.title for item in sheet_stats if item.state != "visible"]
    lines = [
        "# 工作簿只读诊断报告",
        "",
        f"- 源文件：`{source_path.name}`",
        f"- 审计时间：`{datetime.now().astimezone().isoformat(timespec='seconds')}`",
        "- 原件处理：只读打开；未修改、未覆盖。",
        "",
        "## 结论先行",
        "",
        f"1. 实际工作表不是 6 张，而是 **{len(sheet_stats)} 张**："
        + "、".join(f"`{item.title}`" for item in sheet_stats)
        + "。",
        f"2. 隐藏工作表为：{('、'.join(f'`{name}`' for name in hidden) or '无')}。"
        " 此外还存在可见的日期/快照样式工作表 `1.21` 和 `12.23`，这是原指令未列出的重大差异。",
        "3. 四张业务表仍是红人开发表、投产比、红人合作表和红人上线表；`建议`不在数据主链。",
        "4. 日期快照不是四张业务表的完整数据源。它只直接支撑上线表中的部分视频表现列，"
        "并通过上线表/合作表间接进入投产比；CRM、寄样、付款、备注、广告、订单、退款等数据另有来源或人工维护。",
        "5. 不能把所有 `#REF!` 改成 `6.17`：断链公式原本对应不同历史快照日期，强行改向单一快照会篡改历史口径。",
        "",
        "## 工作表盘点",
        "",
        "| 工作表 | 状态 | 声明范围 | 实际非空范围 | 公式 | 字面值 | 仅样式空单元格 | 合并区域 | 数据验证 |",
        "|---|---:|---:|---:|---:|---:|---:|---:|---:|",
    ]
    for item in sheet_stats:
        actual = (
            f"{get_column_letter(item.data_min_column or 1)}{item.data_min_row}:"
            f"{get_column_letter(item.data_max_column or 1)}{item.data_max_row}"
            if item.data_min_row
            else "空"
        )
        lines.append(
            f"| {item.title} | {item.state} | "
            f"A1:{get_column_letter(item.declared_max_column)}{item.declared_max_row} | "
            f"{actual} | {item.formula_cells:,} | {item.literal_cells:,} | "
            f"{item.styled_blank_cells:,} | {item.merged_ranges:,} | {item.data_validations:,} |"
        )

    lines.extend(
        [
            "",
            "## `6.17` 快照复核",
            "",
            f"- 数据行：**{snapshot_summary['data_rows']:,}**。",
            f"- 不同视频链接：**{snapshot_summary['distinct_video_urls']:,}**；"
            f"按“总行数－不同链接数”计算的重复视频链接行：**{snapshot_summary['duplicate_video_url_rows']:,}**。",
            f"- 不同视频 ID：**{snapshot_summary['distinct_video_ids']:,}**。",
            f"- 不同达人 ID：**{snapshot_summary['distinct_creator_ids']:,}**。",
            f"- 完全重复行：**{snapshot_summary['fully_duplicate_rows']:,}**。",
            f"- 发布时间范围：**{snapshot_summary['publish_date_min']} 至 {snapshot_summary['publish_date_max']}**。",
            "- A1 原表头为空，但 A2:A 由内容确认是视频链接。其余 B:P 与 TikTok 视频表现字段相符；"
            "字段盘点已保留原始表头，不把推断写回原件。",
            "- `snapshot_date=2025-06-17` 是导出快照日期；F 列是 `publish_date`，两者已在导入模型中分开。",
            "",
            "## 真实依赖关系",
            "",
            "### 直接来自日期快照",
            "",
            "- 红人上线表的 L/M/N（视频曝光/售出数量/销售额）存在按视频链接查找日期快照 H/I/J 的公式。",
            "- 当前工作簿还出现 `1.21` 与 `12.23` 两张快照式工作表，说明历史快照管理并非只剩 `6.17`，"
            "但各公式对应的日期和口径仍需业务确认。",
            "",
            "### 间接依赖",
            "",
            "- 红人合作表按达人昵称统计红人上线表的上线次数。",
            "- 投产比中的合作量、上线量等公式引用合作表或上线表，因此会间接受日期快照回填影响。",
            "",
            "### 公司人工 CRM",
            "",
            "- 红人开发表主要是开发日期、开发人、红人姓名、邮箱、WhatsApp、归属和备注。",
            "- 红人合作表中的最近联系、开发人、颜色与数量、联系方式、寄样产品、MCN、合作方式、"
            "付款渠道、付款账号、PayPal 和备注以人工维护为主。",
            "- 红人上线表中的登记日期、开发人、AD CODE、投放状态、红人类型和内部备注不能被平台数据覆盖。",
            "",
            "### 其他 TikTok/广告来源",
            "",
            "- 投产比中的订单、店铺浏览、广告花费、广告订单、GMV Max、退款/退货等不能由视频快照完整生成；"
            "它们属于其他 TikTok 报表、广告系统、内部计算或人工值。正式来源仍待业务确认。",
            "",
            "## 公式与断链",
            "",
            f"- 公式总数：**{sum(formula_counts.values()):,}**。",
            f"- 公式文本直接含错误标记的单元格：**{sum(ref_error_counts.values()):,}**。",
            f"- 公式文本或缓存值错误记录：**{len(error_rows):,}**。",
            f"- 外部链接对象：**{external_link_count:,}**。",
            "- 红人上线表断链集中在 L/M/N："
            f"L **{formula_error_column_counts[('Tech-wave红人上线表', 12)]:,}**、"
            f"M **{formula_error_column_counts[('Tech-wave红人上线表', 13)]:,}**、"
            f"N **{formula_error_column_counts[('Tech-wave红人上线表', 14)]:,}**。",
            "",
            "按工作表统计：",
            "",
            "| 工作表 | 公式数 | 含错误标记公式数 |",
            "|---|---:|---:|",
        ]
    )
    for item in sheet_stats:
        lines.append(
            f"| {item.title} | {formula_counts[item.title]:,} | {ref_error_counts[item.title]:,} |"
        )

    lines.extend(
        [
            "",
            "主要跨表公式关系（按公式单元格数）：",
            "",
            "| 来源表 | 目标表 | 公式单元格数 | 目标是否存在 |",
            "|---|---|---:|---:|",
        ]
    )
    for row in relationship_rows[:30]:
        lines.append(
            f"| {row['source_sheet']} | {row['target_sheet']} | "
            f"{row['formula_cell_count']:,} | {row['target_exists']} |"
        )

    lines.extend(
        [
            "",
            "## 为什么不能批量修复 `#REF!`",
            "",
            "断链单元格分布在历史行和经营时间轴中，原本引用的快照日期可能不同。`6.17` 只是一个"
            "导出时点的滚动窗口，不代表所有历史时点。把所有断链改到 `6.17` 会同时混淆快照日期、"
            "视频发布时间、商品维度和历史归因口径。正确做法是保存不可变快照，按稳定 ID、快照日期、"
            "商品维度和明确规则更新自动字段；历史来源无法恢复时应保留为异常，而不是猜测。",
            "",
            "## 审计限制",
            "",
            "- Excel 公式未在 Python 中重算；报告区分“公式文本含错误标记”和“工作簿保存的缓存错误值”。",
            "- 列来源比例是按单元格类型和公式引用进行的审计分类，不代表业务最终所有权已经确认。",
            "- `1.21`、`12.23` 的实际导出日期、店铺、市场、时区和口径仍需用户确认。",
            "",
        ]
    )
    return "\n".join(lines)
