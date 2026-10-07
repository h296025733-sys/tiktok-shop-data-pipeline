from __future__ import annotations

import csv
import hashlib
import json
import math
import os
import shutil
import sqlite3
import tempfile
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from src.config import PROJECT_ROOT
from src.state import STATE_DB, initialize_state_db


NORMALIZED_FIELDS = [
    "source_file",
    "source_sheet",
    "snapshot_date",
    "shop_id",
    "video_url",
    "creator_nickname",
    "creator_id",
    "video_text",
    "video_id",
    "publish_date",
    "product_name_raw",
    "video_views",
    "video_units_sold",
    "video_gmv",
    "gpm",
    "shoppable_video_gmv",
    "video_ctr",
    "traffic_rate",
    "video_completion_rate",
    "video_click_to_order_rate",
    "row_fingerprint",
    "ingested_at",
]

EXPECTED_HEADERS = {
    2: "达人昵称",
    3: "达人ID",
    5: "视频ID",
    6: "发布时间",
    7: "商品",
    8: "视频vv",
    9: "视频商品成交件数",
}


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if math.isnan(value):
            return ""
        if value.is_integer():
            return str(int(value))
        return format(value, ".15g")
    return str(value).strip()


def _number(value: Any) -> float | int | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (int, float)):
        if isinstance(value, float) and math.isnan(value):
            return None
        return value
    text = str(value).strip().replace(",", "").replace("$", "")
    if not text:
        return None
    is_percent = text.endswith("%")
    if is_percent:
        text = text[:-1]
    parsed = float(text)
    if is_percent:
        parsed /= 100
    return int(parsed) if parsed.is_integer() else parsed


def _iso_date(value: Any) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    for parser in (date.fromisoformat,):
        try:
            return parser(text[:10]).isoformat()
        except ValueError:
            pass
    raise ValueError(f"无法解析发布时间：{text!r}")


def _fingerprint(canonical_source_row: dict[str, Any]) -> str:
    payload = json.dumps(
        canonical_source_row,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        default=str,
    )
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _load_existing_ingested_at(path: Path) -> dict[str, str]:
    if not path.exists():
        return {}
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return {
            row["row_fingerprint"]: row["ingested_at"]
            for row in csv.DictReader(handle)
            if row.get("row_fingerprint") and row.get("ingested_at")
        }


def _atomic_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        "w",
        encoding="utf-8-sig",
        newline="",
        delete=False,
        dir=path.parent,
        suffix=".tmp",
    ) as handle:
        writer = csv.DictWriter(handle, fieldnames=NORMALIZED_FIELDS)
        writer.writeheader()
        writer.writerows(rows)
        temp_path = Path(handle.name)
    os.replace(temp_path, path)


def _write_xlsx(path: Path, rows: list[dict[str, Any]], metadata: dict[str, Any]) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "video_performance"
    worksheet.append(NORMALIZED_FIELDS)
    for row in rows:
        worksheet.append([row.get(field) for field in NORMALIZED_FIELDS])

    header_fill = PatternFill("solid", fgColor="0F766E")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    worksheet.row_dimensions[1].height = 36
    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False
    worksheet.auto_filter.ref = worksheet.dimensions

    if worksheet.max_row >= 2:
        table = Table(displayName="VideoPerformance", ref=worksheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        worksheet.add_table(table)

    field_to_column = {
        field: index + 1 for index, field in enumerate(NORMALIZED_FIELDS)
    }
    for field in ("creator_id", "video_id", "shop_id", "row_fingerprint"):
        column = field_to_column[field]
        for row_index in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row_index, column)
            cell.number_format = "@"
            cell.quotePrefix = True
    for field in ("snapshot_date", "publish_date"):
        column = field_to_column[field]
        for row_index in range(2, worksheet.max_row + 1):
            worksheet.cell(row_index, column).number_format = "yyyy-mm-dd"
    for field in ("video_gmv", "gpm", "shoppable_video_gmv"):
        column = field_to_column[field]
        for row_index in range(2, worksheet.max_row + 1):
            worksheet.cell(row_index, column).number_format = '#,##0.00'
    for field in (
        "video_ctr",
        "traffic_rate",
        "video_completion_rate",
        "video_click_to_order_rate",
    ):
        column = field_to_column[field]
        for row_index in range(2, worksheet.max_row + 1):
            worksheet.cell(row_index, column).number_format = "0.00%"
    for field in ("video_views", "video_units_sold"):
        column = field_to_column[field]
        for row_index in range(2, worksheet.max_row + 1):
            worksheet.cell(row_index, column).number_format = "#,##0"

    widths = {
        "source_file": 30,
        "source_sheet": 14,
        "snapshot_date": 14,
        "shop_id": 14,
        "video_url": 45,
        "creator_nickname": 24,
        "creator_id": 22,
        "video_text": 55,
        "video_id": 22,
        "publish_date": 14,
        "product_name_raw": 60,
        "row_fingerprint": 20,
        "ingested_at": 25,
    }
    for field, index in field_to_column.items():
        worksheet.column_dimensions[get_column_letter(index)].width = widths.get(field, 18)

    notes = workbook.create_sheet("说明")
    notes.sheet_view.showGridLines = False
    notes.append(["项目", "值"])
    for key, value in metadata.items():
        notes.append([key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value])
    for cell in notes[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
    notes.column_dimensions["A"].width = 30
    notes.column_dimensions["B"].width = 100
    notes.freeze_panes = "A2"

    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        delete=False, dir=path.parent, suffix=".xlsx"
    ) as handle:
        temp_path = Path(handle.name)
    workbook.save(temp_path)
    os.replace(temp_path, path)


def _write_duplicate_details(path: Path, duplicate_rows: list[dict[str, Any]]) -> None:
    headers = [
        "source_row",
        "duplicate_type",
        "video_url",
        "creator_id",
        "video_id",
        "product_name_raw",
        "row_fingerprint",
    ]
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        writer.writerows(duplicate_rows)


def ingest_snapshot(source_path: Path, sheet_name: str, snapshot_date: date) -> dict[str, Any]:
    source_path = source_path.resolve()
    project_root = PROJECT_ROOT.resolve()
    try:
        source_path.relative_to(project_root)
    except ValueError as exc:
        raise ValueError("导入文件必须位于本项目目录内") from exc

    source_hash = _sha256(source_path)
    snapshot_key = snapshot_date.isoformat()
    workbook = load_workbook(source_path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"工作表不存在：{sheet_name}")
    worksheet = workbook[sheet_name]
    for column_index, expected in EXPECTED_HEADERS.items():
        actual = _text(worksheet.cell(1, column_index).value)
        if actual.lower() != expected.lower():
            raise ValueError(
                f"{sheet_name}!{get_column_letter(column_index)}1 表头不符合预期："
                f"{actual!r} != {expected!r}"
            )

    normalized_dir = (
        PROJECT_ROOT
        / "data"
        / "normalized"
        / "video_performance_snapshots"
        / snapshot_key
    )
    normalized_path = normalized_dir / "video_performance.csv"
    prior_ingested_at = _load_existing_ingested_at(normalized_path)
    current_ingested_at = datetime.now(timezone.utc).isoformat(timespec="seconds")

    rows: list[dict[str, Any]] = []
    source_row_numbers: list[int] = []
    for row_index, values in enumerate(
        worksheet.iter_rows(min_row=2, max_col=16, values_only=True), start=2
    ):
        if not any(value is not None and value != "" for value in values):
            continue
        canonical = {
            "video_url": _text(values[0]),
            "creator_nickname": _text(values[1]),
            "creator_id": _text(values[2]),
            "video_text": _text(values[3]),
            "video_id": _text(values[4]),
            "publish_date": _iso_date(values[5]),
            "product_name_raw": _text(values[6]),
            "video_views": _number(values[7]),
            "video_units_sold": _number(values[8]),
            "video_gmv": _number(values[9]),
            "gpm": _number(values[10]),
            "shoppable_video_gmv": _number(values[11]),
            "video_ctr": _number(values[12]),
            "traffic_rate": _number(values[13]),
            "video_completion_rate": _number(values[14]),
            "video_click_to_order_rate": _number(values[15]),
        }
        fingerprint = _fingerprint(canonical)
        row = {
            "source_file": source_path.name,
            "source_sheet": sheet_name,
            "snapshot_date": snapshot_key,
            "shop_id": "",
            **canonical,
            "row_fingerprint": fingerprint,
            "ingested_at": prior_ingested_at.get(fingerprint, current_ingested_at),
        }
        rows.append(row)
        source_row_numbers.append(row_index)
    workbook.close()

    raw_dir = PROJECT_ROOT / "data" / "raw" / "manual_exports" / snapshot_key
    raw_dir.mkdir(parents=True, exist_ok=True)
    raw_path = raw_dir / f"{source_hash[:12]}_{source_path.name}"
    if raw_path.exists():
        if _sha256(raw_path) != source_hash:
            raise RuntimeError("不可变原始快照路径发生哈希冲突")
    else:
        with tempfile.NamedTemporaryFile(
            delete=False, dir=raw_dir, suffix=source_path.suffix
        ) as handle:
            temp_path = Path(handle.name)
        shutil.copy2(source_path, temp_path)
        if _sha256(temp_path) != source_hash:
            temp_path.unlink(missing_ok=True)
            raise RuntimeError("原始快照复制后哈希校验失败")
        os.replace(temp_path, raw_path)

    url_counts = Counter(row["video_url"] for row in rows if row["video_url"])
    video_counts = Counter(row["video_id"] for row in rows if row["video_id"])
    fingerprint_counts = Counter(row["row_fingerprint"] for row in rows)
    video_products: dict[str, set[str]] = defaultdict(set)
    video_product_counts: Counter[tuple[str, str]] = Counter()
    for row in rows:
        if row["video_id"]:
            video_products[row["video_id"]].add(row["product_name_raw"])
            video_product_counts[(row["video_id"], row["product_name_raw"])] += 1

    duplicate_details: list[dict[str, Any]] = []
    for source_row, row in zip(source_row_numbers, rows, strict=True):
        duplicate_types = []
        if row["video_url"] and url_counts[row["video_url"]] > 1:
            duplicate_types.append("duplicate_video_url")
        if row["video_id"] and video_counts[row["video_id"]] > 1:
            duplicate_types.append("duplicate_video_id")
        if fingerprint_counts[row["row_fingerprint"]] > 1:
            duplicate_types.append("fully_duplicate_row")
        if len(video_products.get(row["video_id"], set())) > 1:
            duplicate_types.append("same_video_different_product")
        if (
            row["video_id"]
            and video_product_counts[(row["video_id"], row["product_name_raw"])] > 1
        ):
            duplicate_types.append("same_video_same_product_multiple_rows")
        if duplicate_types:
            duplicate_details.append(
                {
                    "source_row": source_row,
                    "duplicate_type": "|".join(duplicate_types),
                    "video_url": row["video_url"],
                    "creator_id": row["creator_id"],
                    "video_id": row["video_id"],
                    "product_name_raw": row["product_name_raw"],
                    "row_fingerprint": row["row_fingerprint"],
                }
            )

    numeric_fields = [
        "video_views",
        "video_units_sold",
        "video_gmv",
        "gpm",
        "shoppable_video_gmv",
        "video_ctr",
        "traffic_rate",
        "video_completion_rate",
        "video_click_to_order_rate",
    ]
    rate_fields = [
        "video_ctr",
        "traffic_rate",
        "video_completion_rate",
        "video_click_to_order_rate",
    ]
    negative_cells = [
        {"row": source_row, "field": field, "value": row[field]}
        for source_row, row in zip(source_row_numbers, rows, strict=True)
        for field in numeric_fields
        if isinstance(row[field], (int, float)) and row[field] < 0
    ]
    out_of_range_rates = [
        {"row": source_row, "field": field, "value": row[field]}
        for source_row, row in zip(source_row_numbers, rows, strict=True)
        for field in rate_fields
        if isinstance(row[field], (int, float)) and not 0 <= row[field] <= 1
    ]
    invalid_numeric_types = [
        {"row": source_row, "field": field, "value": repr(row[field])}
        for source_row, row in zip(source_row_numbers, rows, strict=True)
        for field in numeric_fields
        if row[field] is not None and not isinstance(row[field], (int, float))
    ]
    zero_gmv_positive_units = [
        source_row
        for source_row, row in zip(source_row_numbers, rows, strict=True)
        if (row["video_units_sold"] or 0) > 0 and (row["video_gmv"] or 0) == 0
    ]
    positive_gmv_zero_units = [
        source_row
        for source_row, row in zip(source_row_numbers, rows, strict=True)
        if (row["video_gmv"] or 0) > 0 and (row["video_units_sold"] or 0) == 0
    ]
    publish_dates = [row["publish_date"] for row in rows if row["publish_date"]]
    suspicious_ids = [
        {"row": source_row, "field": field, "value": row[field]}
        for source_row, row in zip(source_row_numbers, rows, strict=True)
        for field in ("creator_id", "video_id")
        if row[field] and ("e+" in row[field].lower() or "e-" in row[field].lower())
    ]
    quality = {
        "source_file": source_path.name,
        "source_sha256": source_hash,
        "source_sheet": sheet_name,
        "snapshot_date": snapshot_key,
        "shop_id": None,
        "row_count": len(rows),
        "empty_video_id": sum(not row["video_id"] for row in rows),
        "empty_creator_id": sum(not row["creator_id"] for row in rows),
        "empty_video_url": sum(not row["video_url"] for row in rows),
        "distinct_video_urls": len(url_counts),
        "duplicate_video_url_rows": len(rows) - len(url_counts),
        "distinct_video_ids": len(video_counts),
        "duplicate_video_id_rows": len(rows) - len(video_counts),
        "fully_duplicate_rows": sum(count - 1 for count in fingerprint_counts.values()),
        "same_video_different_product_video_count": sum(
            len(products) > 1 for products in video_products.values()
        ),
        "same_video_same_product_extra_rows": sum(
            count - 1 for count in video_product_counts.values() if count > 1
        ),
        "invalid_numeric_type_count": len(invalid_numeric_types),
        "negative_numeric_count": len(negative_cells),
        "percentage_out_of_range_count": len(out_of_range_rates),
        "positive_units_zero_gmv_rows": len(zero_gmv_positive_units),
        "positive_gmv_zero_units_rows": len(positive_gmv_zero_units),
        "publish_date_min": min(publish_dates) if publish_dates else None,
        "publish_date_max": max(publish_dates) if publish_dates else None,
        "scientific_notation_id_count": len(suspicious_ids),
        "examples": {
            "invalid_numeric_types": invalid_numeric_types[:20],
            "negative_numeric": negative_cells[:20],
            "percentage_out_of_range": out_of_range_rates[:20],
            "positive_units_zero_gmv_source_rows": zero_gmv_positive_units[:50],
            "positive_gmv_zero_units_source_rows": positive_gmv_zero_units[:50],
            "scientific_notation_ids": suspicious_ids[:20],
        },
    }

    _atomic_csv(normalized_path, rows)
    export_csv = PROJECT_ROOT / "exports" / "csv" / snapshot_key / "video_performance.csv"
    _atomic_csv(export_csv, rows)
    export_xlsx = (
        PROJECT_ROOT / "exports" / "xlsx" / snapshot_key / "video_performance.xlsx"
    )
    _write_xlsx(
        export_xlsx,
        rows,
        {
            "source_file": source_path.name,
            "source_sha256": source_hash,
            "source_sheet": sheet_name,
            "snapshot_date": snapshot_key,
            "shop_id": "未知（未编造）",
            "row_count": len(rows),
            "说明": "原始记录未去重；snapshot_date 与 publish_date 已分离。",
        },
    )

    report_dir = PROJECT_ROOT / "exports" / "reports" / snapshot_key
    report_dir.mkdir(parents=True, exist_ok=True)
    quality_path = report_dir / "data_quality_report.json"
    quality_path.write_text(
        json.dumps(quality, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    duplicate_path = report_dir / "duplicate_rows.csv"
    _write_duplicate_details(duplicate_path, duplicate_details)
    candidate_rules_path = report_dir / "candidate_deduplication_rules.json"
    candidate_rules_path.write_text(
        json.dumps(
            {
                "status": "候选规则，未经用户确认不得删除记录",
                "recommended_snapshot_key": [
                    "shop_id",
                    "snapshot_date",
                    "video_id",
                    "product_id",
                ],
                "fallback_when_product_id_missing": [
                    "shop_id",
                    "snapshot_date",
                    "video_id",
                    "normalized_product_label",
                    "row_fingerprint",
                ],
                "current_file_limitations": [
                    "shop_id 未知",
                    "product_id 不存在",
                    "商品只有显示名称",
                ],
                "do_not_use_alone": ["video_id", "video_url"],
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    raw_metadata_path = raw_dir / f"{source_hash[:12]}_metadata.json"
    if not raw_metadata_path.exists():
        raw_metadata_path.write_text(
            json.dumps(
                {
                    "source_file": source_path.name,
                    "source_sha256": source_hash,
                    "source_sheet": sheet_name,
                    "snapshot_date": snapshot_key,
                    "shop_id": None,
                    "raw_path": str(raw_path),
                },
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )

    initialize_state_db(STATE_DB)
    with sqlite3.connect(STATE_DB) as connection:
        connection.execute(
            """
            INSERT INTO import_snapshots (
                source_sha256, source_sheet, snapshot_date, ingested_at,
                normalized_path, row_count
            ) VALUES (?, ?, ?, ?, ?, ?)
            ON CONFLICT(source_sha256, source_sheet, snapshot_date) DO UPDATE SET
                normalized_path = excluded.normalized_path,
                row_count = excluded.row_count
            """,
            (
                source_hash,
                sheet_name,
                snapshot_key,
                current_ingested_at,
                str(normalized_path),
                len(rows),
            ),
        )

    return {
        "row_count": len(rows),
        "source_sha256": source_hash,
        "raw_snapshot": str(raw_path),
        "normalized_csv": str(normalized_path),
        "export_csv": str(export_csv),
        "export_xlsx": str(export_xlsx),
        "quality_report": str(quality_path),
        "duplicate_details": str(duplicate_path),
        "candidate_deduplication_rules": str(candidate_rules_path),
    }
