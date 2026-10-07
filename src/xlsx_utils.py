from __future__ import annotations

from collections.abc import Iterable
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


HEADER_FILL = PatternFill("solid", fgColor="0F766E")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SUBTLE_FILL = PatternFill("solid", fgColor="E6FFFB")


def style_tabular_sheet(
    worksheet,
    *,
    table_name: str | None = None,
    freeze_rows: int = 1,
    max_width: int = 42,
) -> None:
    worksheet.freeze_panes = f"A{freeze_rows + 1}"
    worksheet.sheet_view.showGridLines = False
    if worksheet.max_row >= 1 and worksheet.max_column >= 1:
        worksheet.row_dimensions[1].height = 36
        for cell in worksheet[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        worksheet.auto_filter.ref = worksheet.dimensions
        if table_name and worksheet.max_row >= 2:
            table = Table(displayName=table_name, ref=worksheet.dimensions)
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            worksheet.add_table(table)

    for column_index in range(1, worksheet.max_column + 1):
        values = [
            str(worksheet.cell(row_index, column_index).value or "")
            for row_index in range(1, min(worksheet.max_row, 250) + 1)
        ]
        width = min(max(max((len(value) for value in values), default=8) + 2, 10), max_width)
        worksheet.column_dimensions[get_column_letter(column_index)].width = width


def write_rows_to_workbook(
    path: Path,
    *,
    sheet_name: str,
    headers: list[str],
    rows: Iterable[Iterable[object]],
    table_name: str,
) -> Path:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    worksheet.append(headers)
    for row in rows:
        values = list(row)
        worksheet.append(values)
        row_index = worksheet.max_row
        for column_index, value in enumerate(values, start=1):
            if isinstance(value, str) and value.startswith("="):
                worksheet.cell(row_index, column_index).data_type = "s"
    style_tabular_sheet(worksheet, table_name=table_name)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path
