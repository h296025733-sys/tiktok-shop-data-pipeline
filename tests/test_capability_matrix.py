from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from src import capability_matrix


def test_capability_matrix_marks_real_tests_consistently(
    tmp_path: Path, monkeypatch
) -> None:
    project = tmp_path / "project"
    config_dir = project / "config" / "datasets"
    config_dir.mkdir(parents=True)
    source_config = (
        capability_matrix.PROJECT_ROOT
        / "config"
        / "datasets"
        / "tiktok_capabilities.json"
    )
    (config_dir / "tiktok_capabilities.json").write_text(
        source_config.read_text(encoding="utf-8"), encoding="utf-8"
    )
    monkeypatch.setattr(capability_matrix, "PROJECT_ROOT", project)
    output = capability_matrix.build_capability_matrix()
    workbook = load_workbook(output, read_only=True, data_only=True)
    worksheet = workbook["API能力矩阵"]
    statuses = {row[18] for row in worksheet.iter_rows(min_row=2, values_only=True)}
    endpoints = {row[3] for row in worksheet.iter_rows(min_row=2, values_only=True)}
    assert statuses <= {
        "尚未验证", "已由官方文档确认", "当前应用已配置",
        "已由真实店铺测试成功", "需要额外权限", "已确认不可用",
    }
    assert "尚未验证" in statuses
    assert "已由真实店铺测试成功" in statuses
    tested_rows = [row for row in worksheet.iter_rows(min_row=2, values_only=True) if row[18] == "已由真实店铺测试成功"]
    assert tested_rows
    assert all(row[6] == "是" and row[7] == "是" and row[8] == "成功" for row in tested_rows)
    assert endpoints <= {
        None,
        "/authorization/202309/shops",
        "/analytics/202409/shop_videos/{video_id}/products/performance",
    }
    assert "/authorization/202309/shops" in endpoints
    assert "/analytics/202409/shop_videos/{video_id}/products/performance" in endpoints
